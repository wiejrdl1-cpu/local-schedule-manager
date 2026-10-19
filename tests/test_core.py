from __future__ import annotations

import sqlite3
import json
import os
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app import ScheduleApp, create_portable_backup_archive
from database import Database, RELEASE_APP_DIR_NAME, default_data_dir
from excel_import import (
    TASK_TEMPLATE_EXAMPLE_TITLE,
    TASK_TEMPLATE_HEADERS,
    apply_import,
    apply_task_import,
    create_task_template,
    preview_import,
    preview_task_import,
)


class CoreFlowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = Database(self.root / "schedule.db")
        self.db.set_initial_password("test-password-123")

    def tearDown(self) -> None:
        self.db.close()
        self.temp.cleanup()

    def test_packaged_app_uses_separate_empty_data_directory(self) -> None:
        with patch.dict(os.environ, {"LOCALAPPDATA": str(self.root)}, clear=False):
            os.environ.pop("DDAY_MANAGER_DATA_DIR", None)
            with patch.object(sys, "frozen", True, create=True):
                self.assertEqual(default_data_dir(), self.root / RELEASE_APP_DIR_NAME)

    def test_packaged_app_prefers_portable_data_beside_executable(self) -> None:
        executable = self.root / "Task It Easy.exe"
        portable = self.root / "data"
        portable.mkdir()
        (portable / "portable.marker").write_text("portable", encoding="utf-8")
        with patch.dict(os.environ, {"LOCALAPPDATA": str(self.root / "local")}, clear=False):
            os.environ.pop("DDAY_MANAGER_DATA_DIR", None)
            with patch.object(sys, "frozen", True, create=True), patch.object(sys, "executable", str(executable)):
                self.assertEqual(default_data_dir(), portable)

    def _write_excel(
        self,
        deadline: str = "2026-09-09 18:00",
        stage: str = "접수",
        received_at: str = "2026-08-10",
        applicant_name: str = "가상민원인",
    ) -> Path:
        path = self.root / "받은민원.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["순번", "접수번호", "민원명", "처리단계", "받은일자", "접수일자", "처리기한", "민원인", "접수구분"])
        ws.append([1, "R-2026-001", "맞춤형복지급여", stage, received_at, "2026-08-11", deadline, applicant_name, "연계"])
        wb.save(path)
        return path

    def test_import_duplicate_and_deadline_change_history(self) -> None:
        path = self._write_excel()
        first = preview_import(path, self.db)
        self.assertEqual([item.status for item in first], ["신규"])
        counts = apply_import(path, self.db, first)
        self.assertEqual(counts["new"], 1)

        duplicate = preview_import(path, self.db)
        self.assertEqual([item.status for item in duplicate], ["중복"])

        self._write_excel(
            "2026-10-09 18:00",
            stage="처리완료",
            received_at="2026-08-30",
            applicant_name="엑셀에서변경된이름",
        )
        changed = preview_import(path, self.db)
        self.assertEqual([item.status for item in changed], ["변경"])
        apply_import(path, self.db, changed)
        complaint = self.db.list_complaints()[0]
        self.assertEqual(complaint["original_deadline"], "2026-09-09 18:00")
        self.assertEqual(complaint["current_deadline"], "2026-10-09 18:00")
        self.assertEqual(complaint["processing_stage"], "접수")
        self.assertEqual(complaint["received_at"], "2026-08-10")
        self.assertEqual(complaint["applicant_name"], "가상민원인")
        self.assertEqual(len(self.db.extension_history(complaint["id"])), 1)

    def test_import_completes_active_complaints_missing_from_new_workbook(self) -> None:
        self.db.add_complaint(
            {
                "receipt_no": "R-OLD-001",
                "complaint_name": "이전 엑셀에만 있던 민원",
                "received_at": "2026-08-01",
                "deadline": "2026-08-20",
            }
        )
        path = self._write_excel()
        items = preview_import(path, self.db)
        counts = apply_import(path, self.db, items)
        self.assertEqual(counts["completed"], 1)
        completed = [
            item
            for item in self.db.list_complaints(include_completed=True)
            if item["receipt_no"] == "R-OLD-001"
        ]
        self.assertEqual(completed[0]["status"], "완료")

    def test_import_with_errors_does_not_auto_complete_existing_complaints(self) -> None:
        self.db.add_complaint(
            {
                "receipt_no": "R-KEEP-001",
                "complaint_name": "유지되어야 하는 민원",
                "received_at": "2026-08-01",
                "deadline": "2026-08-20",
            }
        )
        path = self.root / "오류민원.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["접수번호", "민원명", "받은일자", "처리기한"])
        sheet.append(["R-ERROR-001", "", "2026-08-10", "2026-08-20"])
        workbook.save(path)
        items = preview_import(path, self.db)
        counts = apply_import(path, self.db, items)
        self.assertEqual(counts["error"], 1)
        self.assertEqual(counts["completed"], 0)
        self.assertEqual(self.db.complaint_by_receipt_no("R-KEEP-001")["status"], "진행")

    def test_sensitive_values_are_not_plaintext_in_database(self) -> None:
        self.db.add_task(
            {
                "title": "서류 확인",
                "applicant_name": "가상민원인",
                "birth_date": "1950-03-12",
                "registered_at": "2026-08-15",
                "deadline": "2026-08-20",
                "content": "가상의 민감 메모",
                "priority": "높음",
            }
        )
        raw_conn = sqlite3.connect(self.db.path)
        try:
            raw = raw_conn.execute(
                "SELECT applicant_name_enc, birth_date_enc, content_enc FROM tasks"
            ).fetchone()
        finally:
            raw_conn.close()
        joined = " ".join(raw)
        self.assertNotIn("가상민원인", joined)
        self.assertNotIn("1950-03-12", joined)
        self.assertNotIn("가상의 민감 메모", joined)
        self.assertEqual(self.db.list_tasks()[0]["applicant_name"], "가상민원인")

    def test_backup_is_created(self) -> None:
        backup = self.db.backup()
        self.assertTrue(backup.exists())
        self.assertGreater(backup.stat().st_size, 0)

    def test_portable_backup_contains_executable_database_and_attachments(self) -> None:
        executable = self.root / "Task It Easy.exe"
        executable.write_bytes(b"sample-executable")
        attachment = self.root / "sample.txt"
        attachment.write_text("attachment", encoding="utf-8")
        task_id = self.db.add_task(
            {"title": "백업 시험", "registered_at": "2026-10-15", "deadline": "2026-10-20"}
        )
        self.db.add_attachment("task", task_id, attachment)
        target = self.root / "portable.zip"
        create_portable_backup_archive(self.db, executable, target)
        with zipfile.ZipFile(target) as archive:
            names = set(archive.namelist())
        self.assertIn("Task It Easy.exe", names)
        self.assertIn("data/schedule.db", names)
        self.assertIn("data/portable.marker", names)
        self.assertTrue(any(name.startswith("data/attachments/") for name in names))

    def test_notification_schedules_and_marks_triggers(self) -> None:
        reference = datetime(2026, 10, 15, 9, 0, 0)
        with patch("database.datetime") as mocked_datetime:
            mocked_datetime.now.return_value = reference
            mocked_datetime.strptime = datetime.strptime
            notification_id = self.db.add_notification(
                {
                    "title": "서류 확인",
                    "message": "민원 서류를 확인하세요.",
                    "schedule_type": "interval",
                    "interval_minutes": 30,
                    "enabled": True,
                }
            )
        item = self.db.notification_by_id(notification_id)
        self.assertEqual(item["message"], "민원 서류를 확인하세요.")
        self.assertEqual(item["next_trigger_at"], "2026-10-15T09:30:00")
        due_time = datetime(2026, 10, 15, 9, 30, 0)
        self.assertEqual([value["id"] for value in self.db.due_notifications(due_time)], [notification_id])
        self.db.mark_notification_triggered(notification_id, due_time)
        self.assertEqual(self.db.notification_by_id(notification_id)["next_trigger_at"], "2026-10-15T10:00:00")

    def test_fixed_overtime_notifications_are_seeded_once_and_cannot_be_deleted(self) -> None:
        self.db.ensure_default_notifications()
        self.db.ensure_default_notifications()
        items = self.db.list_notifications()
        fixed = [item for item in items if item["fixed_key"]]
        self.assertEqual({item["title"] for item in fixed}, {"평일 초과근무", "주말 초과근무"})
        self.assertEqual({item["message"] for item in fixed}, {"초과 근무를 입력하세요"})
        self.assertEqual(len(fixed), 2)
        with self.assertRaisesRegex(ValueError, "삭제할 수 없습니다"):
            self.db.delete_notification(int(fixed[0]["id"]))

        customized = fixed[0]
        self.db.update_notification(
            int(customized["id"]),
            {
                **customized,
                "message": "사용자가 지정한 알림 문구",
                "enabled": True,
            },
        )
        self.db.ensure_default_notifications()
        self.assertEqual(
            self.db.notification_by_id(int(customized["id"]))["message"],
            "사용자가 지정한 알림 문구",
        )

    def test_fixed_overtime_notification_time_windows(self) -> None:
        weekday_reference = datetime(2026, 10, 16, 9, 20, 0)  # Friday
        weekday_next = Database._next_notification_at(
            "weekday_overtime", notify_time="15", reference=weekday_reference
        )
        self.assertEqual(weekday_next, datetime(2026, 10, 16, 18, 15, 0))
        weekday_late_reference = datetime(2026, 10, 16, 23, 20, 0)  # Friday
        weekday_after_weekend = Database._next_notification_at(
            "weekday_overtime", notify_time="15", reference=weekday_late_reference
        )
        self.assertEqual(weekday_after_weekend, datetime(2026, 10, 19, 0, 15, 0))
        weekend_reference = datetime(2026, 10, 17, 23, 20, 0)  # Saturday
        weekend_next = Database._next_notification_at(
            "weekend_overtime", notify_time="15", reference=weekend_reference
        )
        self.assertEqual(weekend_next, datetime(2026, 10, 18, 0, 15, 0))

    def test_processing_stage_change_is_recorded(self) -> None:
        complaint_id = self.db.add_complaint(
            {
                "receipt_no": "R-STAGE-001",
                "complaint_name": "단계 기록 시험",
                "processing_stage": "1. 공적자료요청",
                "received_at": "2026-08-15",
                "registered_at": "2026-08-15",
                "deadline": "2026-09-15",
                "applicant_name": "가상민원인",
            }
        )
        self.db.update_processing_stage(complaint_id, "3. 조사중 > 3-2. 가정방문")
        complaint = self.db.complaint_by_id(complaint_id)
        self.assertEqual(complaint["processing_stage"], "3. 조사중 > 3-2. 가정방문")
        history = self.db.processing_stage_history(complaint_id)
        self.assertEqual(history[0]["previous_stage"], "1. 공적자료요청")
        self.assertEqual(history[0]["new_stage"], "3. 조사중 > 3-2. 가정방문")

    def test_duplicate_receipt_number_has_friendly_message(self) -> None:
        data = {
            "receipt_no": "R-DUPLICATE-001",
            "complaint_name": "중복 확인",
            "processing_stage": "1. 공적자료요청",
            "received_at": "2026-08-15",
            "deadline": "2026-09-15",
        }
        self.db.add_complaint(data)
        with self.assertRaisesRegex(ValueError, "이미 등록되어 있습니다"):
            self.db.add_complaint(data)

    def test_auto_unlock_uses_current_windows_user(self) -> None:
        self.db.enable_auto_unlock("test-password-123")
        self.db.protector = None
        self.assertTrue(self.db.try_auto_unlock())
        self.assertIsNotNone(self.db.protector)

    def test_complaint_and_task_can_be_fully_edited(self) -> None:
        complaint_id = self.db.add_complaint(
            {
                "receipt_no": "R-EDIT-OLD",
                "complaint_name": "수정 전 민원",
                "processing_stage": "1. 공적자료요청",
                "received_at": "2026-08-15",
                "registered_at": "2026-08-15",
                "deadline": "2026-09-15",
                "applicant_name": "수정 전 이름",
            }
        )
        self.db.update_complaint(
            complaint_id,
            {
                "receipt_no": "R-EDIT-NEW",
                "complaint_name": "수정 후 민원",
                "processing_stage": "3. 조사중 > 3-1. 소득조사",
                "received_at": "2026-08-16",
                "registered_at": "2026-08-16",
                "deadline": "2026-10-15",
                "applicant_name": "수정 후 이름",
                "birth_date": "1991-11-01",
                "intake_type": "수기",
                "assignee": "담당자",
                "memo": "수정 메모",
            },
        )
        complaint = self.db.complaint_by_id(complaint_id)
        self.assertEqual(complaint["receipt_no"], "R-EDIT-NEW")
        self.assertEqual(complaint["applicant_name"], "수정 후 이름")
        self.assertEqual(complaint["current_deadline"], "2026-10-15")
        self.assertEqual(complaint["original_deadline"], "2026-09-15")

        task_id = self.db.add_task(
            {"title": "수정 전", "registered_at": "2026-08-15", "deadline": "2026-08-20", "content": "전"}
        )
        self.db.update_task(
            task_id,
            {"title": "수정 후", "registered_at": "2026-08-16", "deadline": "2026-08-25", "priority": "긴급", "content": "후"},
        )
        task = self.db.task_by_id(task_id)
        self.assertEqual(task["title"], "수정 후")
        self.assertEqual(task["priority"], "긴급")

    def test_progress_state_is_saved_with_current_and_completed_nodes(self) -> None:
        complaint_id = self.db.add_complaint(
            {
                "receipt_no": "R-PROGRESS-001",
                "complaint_name": "진행 현황 시험",
                "processing_stage": "1. 신청",
                "received_at": "2026-08-15",
                "deadline": "2026-09-15",
            }
        )
        state = {
            "current": "b:investigation:income-check",
            "completed": ["s:application", "s:public-request"],
        }
        self.db.update_progress_state(
            complaint_id,
            state,
            "3. 조사중 > 3-1. 소득조사",
        )
        complaint = self.db.complaint_by_id(complaint_id)
        self.assertEqual(complaint["processing_stage"], "3. 조사중 > 3-1. 소득조사")
        self.assertIn("s:application", complaint["progress_state"])

    def test_completed_items_can_be_restored(self) -> None:
        complaint_id = self.db.add_complaint(
            {"receipt_no": "R-RESTORE", "complaint_name": "복원 시험", "received_at": "2026-08-15", "deadline": "2026-09-15"}
        )
        self.db.mark_complete("complaint", complaint_id)
        self.assertEqual(self.db.complaint_by_id(complaint_id)["status"], "완료")
        self.db.restore_completed("complaint", complaint_id)
        restored = self.db.complaint_by_id(complaint_id)
        self.assertEqual(restored["status"], "진행")
        self.assertEqual(restored["completed_at"], "")

    def test_task_progress_stages_are_saved(self) -> None:
        task_id = self.db.add_task(
            {"title": "수시 단계 시험", "registered_at": "2026-08-15", "deadline": "2026-09-15", "processing_stage": "접수"}
        )
        self.db.update_task_progress_state(
            task_id,
            {"current": "s:task-working", "completed": ["s:task-received"]},
            "작업중",
        )
        task = self.db.task_by_id(task_id)
        self.assertEqual(task["processing_stage"], "작업중")
        self.assertIn("task-received", task["progress_state"])

    def test_task_excel_template_and_blank_values_are_preserved(self) -> None:
        template = self.root / "수시업무_샘플.xlsx"
        create_task_template(template)
        workbook = load_workbook(template)
        self.assertEqual([cell.value for cell in workbook.active[1]], TASK_TEMPLATE_HEADERS)
        example = [cell.value for cell in workbook.active[2]]
        self.assertEqual(example[0], TASK_TEMPLATE_EXAMPLE_TITLE)
        self.assertEqual(example[1:3], ["2026-08-29", "2026-09-05"])
        self.assertIn("YYYY-MM-DD", example[3])
        self.assertEqual(preview_task_import(template), [])

        path = self.root / "수시업무.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(TASK_TEMPLATE_HEADERS)
        sheet.append(["공백 확인", "2026-08-20", None, "  앞뒤 공백 유지  "])
        sheet.append([None, None, "2026-09-01", None])
        workbook.save(path)

        items = preview_task_import(path)
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0].data["deadline"], "")
        self.assertEqual(items[0].data["priority"], "  앞뒤 공백 유지  ")
        self.assertEqual(items[1].data["title"], "")
        counts = apply_task_import(path, self.db, items, "접수")
        self.assertEqual(counts["new"], 2)
        tasks = self.db.list_tasks()
        self.assertEqual(tasks[0]["priority"], "  앞뒤 공백 유지  ")
        self.assertEqual(tasks[0]["deadline"], "")
        self.assertEqual(tasks[1]["title"], "")

    def test_calendar_event_details_are_saved_and_encrypted(self) -> None:
        event_id = self.db.add_calendar_event("2026-08-29", "회의 일정", "가상의 상세 회의 내용", "14:30", "부서일정")
        event = self.db.list_calendar_events()[0]
        self.assertEqual(event["id"], event_id)
        self.assertEqual(event["event_date"], "2026-08-29")
        self.assertEqual(event["event_time"], "14:30")
        self.assertEqual(event["schedule_type"], "부서일정")
        self.assertEqual(event["title"], "회의 일정")
        self.assertEqual(event["details"], "가상의 상세 회의 내용")

        raw_conn = sqlite3.connect(self.db.path)
        try:
            encrypted_details = raw_conn.execute(
                "SELECT details_enc FROM calendar_events WHERE id=?", (event_id,)
            ).fetchone()[0]
        finally:
            raw_conn.close()
        self.assertNotIn("가상의 상세 회의 내용", encrypted_details)

        self.db.update_calendar_event(
            event_id,
            "2026-08-30",
            "수정된 회의 일정",
            "수정된 상세 내용",
            "18:00",
            "개인일정",
        )
        updated = self.db.calendar_event_by_id(event_id)
        self.assertEqual(updated["event_date"], "2026-08-30")
        self.assertEqual(updated["event_time"], "18:00")
        self.assertEqual(updated["schedule_type"], "개인일정")
        self.assertEqual(updated["title"], "수정된 회의 일정")
        self.assertEqual(updated["details"], "수정된 상세 내용")

        self.db.delete_calendar_event(event_id)
        self.assertIsNone(self.db.calendar_event_by_id(event_id))

    def test_daily_weekly_and_monthly_recurring_events_are_expanded(self) -> None:
        daily_id = self.db.add_calendar_event(
            "2026-01-01",
            "격일 일정",
            recurrence={"repeat_type": "daily", "repeat_interval": 2, "repeat_end_date": "2026-01-07"},
        )
        weekly_id = self.db.add_calendar_event(
            "2026-01-05",
            "격주 월수 일정",
            recurrence={
                "repeat_type": "weekly",
                "repeat_interval": 2,
                "repeat_weekdays": "0,2",
                "repeat_end_date": "2026-02-02",
            },
        )
        monthly_id = self.db.add_calendar_event(
            "2026-01-30",
            "매월 31일 일정",
            recurrence={
                "repeat_type": "monthly",
                "repeat_interval": 1,
                "repeat_month_day": 31,
                "repeat_end_date": "2026-04-30",
            },
        )
        occurrences = self.db.calendar_events_between(date(2026, 1, 1), date(2026, 4, 30))
        by_id: dict[int, list[str]] = {}
        for event in occurrences:
            by_id.setdefault(int(event["id"]), []).append(str(event["event_date"]))
        self.assertEqual(by_id[daily_id], ["2026-01-01", "2026-01-03", "2026-01-05", "2026-01-07"])
        self.assertEqual(
            by_id[weekly_id],
            ["2026-01-05", "2026-01-07", "2026-01-19", "2026-01-21", "2026-02-02"],
        )
        self.assertEqual(by_id[monthly_id], ["2026-01-31", "2026-03-31"])

    def test_attachments_are_copied_for_all_item_types_and_cleaned_up(self) -> None:
        source = self.root / "업무 참고자료.txt"
        source.write_text("가상의 첨부파일 내용", encoding="utf-8")
        complaint_id = self.db.add_complaint(
            {"receipt_no": "R-ATTACH", "complaint_name": "첨부 민원", "received_at": "2026-01-01", "deadline": "2026-01-10"}
        )
        task_id = self.db.add_task(
            {"title": "첨부 업무", "registered_at": "2026-01-01", "deadline": "2026-01-10"}
        )
        event_id = self.db.add_calendar_event("2026-01-05", "첨부 일정")
        for kind, item_id in (("complaint", complaint_id), ("task", task_id), ("calendar_event", event_id)):
            attachment_id = self.db.add_attachment(kind, item_id, source)
            attachment = self.db.attachment_by_id(attachment_id)
            self.assertEqual(attachment["display_name"], source.name)
            self.assertEqual(Path(attachment["path"]).read_text(encoding="utf-8"), "가상의 첨부파일 내용")
            self.assertEqual(len(self.db.list_attachments(kind, item_id)), 1)
        event_attachment_path = Path(self.db.list_attachments("calendar_event", event_id)[0]["path"])
        self.db.delete_calendar_event(event_id)
        self.assertFalse(event_attachment_path.exists())
        self.assertEqual(self.db.list_attachments("calendar_event", event_id), [])

    def test_existing_calendar_events_default_to_personal_schedule(self) -> None:
        legacy = sqlite3.connect(self.root / "legacy.db")
        legacy.execute(
            "CREATE TABLE calendar_events ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, event_date TEXT NOT NULL, "
            "event_time TEXT NOT NULL DEFAULT '', title TEXT NOT NULL, "
            "details_enc TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL)"
        )
        legacy.execute(
            "INSERT INTO calendar_events(event_date, title, created_at, updated_at) VALUES('2026-09-01', '기존 일정', '', '')"
        )
        legacy.commit()
        legacy.close()

        migrated = Database(self.root / "legacy.db")
        try:
            migrated.set_initial_password("migration-test-password")
            migrated_event = migrated.list_calendar_events()[0]
            self.assertEqual(migrated_event["schedule_type"], "개인일정")
            self.assertEqual(migrated_event["repeat_type"], "none")
        finally:
            migrated.close()

    def test_task_progress_supports_detail_branches(self) -> None:
        app = object.__new__(ScheduleApp)
        app.task_stage_config = [
            {"id": "received", "name": "접수", "branches": []},
            {
                "id": "working",
                "name": "처리중",
                "branches": [
                    {"id": "review", "name": "검토"},
                    {"id": "approval", "name": "결재"},
                ],
            },
        ]
        task = {
            "stage_config": "",
            "processing_stage": "처리중 > 결재",
            "progress_state": json.dumps(
                {"current": "b:working:approval", "completed": ["s:received", "b:working:review"]},
                ensure_ascii=False,
            ),
        }
        self.assertEqual(app._progress_node_keys(app.task_stage_config), ["s:received", "b:working:review", "b:working:approval"])
        self.assertEqual(app._task_progress_state(task)["current"], "b:working:approval")
        self.assertEqual(app._task_stage_text("b:working:approval"), "처리중 > 결재")

        app.db = self.db
        app.refresh_all = lambda: None
        task_id = self.db.add_task(
            {"title": "분기 처리", "registered_at": "2026-08-20", "deadline": "2026-09-10", "processing_stage": "접수"}
        )
        app._toggle_task_progress(task_id, "b:working:review", True)
        updated = self.db.task_by_id(task_id)
        self.assertEqual(updated["processing_stage"], "처리중 > 결재")
        self.assertIn("b:working:review", updated["progress_state"])

    def test_dashboard_deadline_filters_match_card_counts(self) -> None:
        today = date.today()
        self.assertTrue(ScheduleApp._matches_dashboard_filter(today.isoformat(), "오늘 마감"))
        self.assertFalse(ScheduleApp._matches_dashboard_filter(today.isoformat(), "3일 이내"))
        self.assertTrue(
            ScheduleApp._matches_dashboard_filter((today + timedelta(days=3)).isoformat(), "3일 이내")
        )
        self.assertFalse(
            ScheduleApp._matches_dashboard_filter((today + timedelta(days=4)).isoformat(), "3일 이내")
        )
        self.assertTrue(
            ScheduleApp._matches_dashboard_filter((today + timedelta(days=4)).isoformat(), "7일 이내")
        )
        self.assertTrue(
            ScheduleApp._matches_dashboard_filter((today + timedelta(days=7)).isoformat(), "7일 이내")
        )
        self.assertFalse(
            ScheduleApp._matches_dashboard_filter((today + timedelta(days=3)).isoformat(), "7일 이내")
        )
        self.assertTrue(
            ScheduleApp._matches_dashboard_filter((today - timedelta(days=1)).isoformat(), "기한 초과")
        )


if __name__ == "__main__":
    unittest.main()
